
"""
Professional PowerPoint Generator.

Pipeline:
1. Convert Pydantic presentation data into a dictionary.
2. Render slides with the chosen engine: "compono", "python-pptx" or
   "both" (Compono layout + python-pptx brand pass on the same file).
3. Validate the final PPTX with Office-Kit (mandatory for every engine;
   a failed validation raises PptxValidationError and keeps the file).
4. Check layout (off-slide shapes, overlaps, tiny fonts) and images.
5. Return the generated file path.
"""

import time
from pathlib import Path
from typing import Any

from pptx import Presentation as PptPresentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from config import PPTX_ENGINE
from office.officekit_checks import print_review, review_document
from office.officekit_utils import render_rich_pptx
from utils.perf import PERF


def _to_dict(value: Any) -> dict:
    """Convert a Pydantic model or dictionary into a dictionary."""

    if hasattr(value, "model_dump"):
        data = value.model_dump()
    elif hasattr(value, "dict"):
        data = value.dict()
    elif isinstance(value, dict):
        data = value
    else:
        raise TypeError(
            "Presentation must be a Pydantic model or dictionary."
        )

    if not isinstance(data, dict):
        raise TypeError("Converted presentation data must be a dictionary.")

    return data


def _count_picture_shapes(pptx_file: Path) -> int:
    """Return the number of embedded picture shapes in a PPTX file."""
    return check_pptx_layout(pptx_file)["pictures"]


# Content shapes that must never overlap each other.
_CONTENT_SHAPES = {
    "Slide Title", "Slide Subtitle", "Body Paragraph", "Bullet List",
    "Content Table", "Slide Image", "Image Caption", "Card", "Chart",
    "Process Step",
}


def check_pptx_layout(pptx_file: Path) -> dict:
    """
    Single pass over the saved PPTX: count embedded pictures and report
    off-slide shapes, overlapping content shapes and tiny fonts.
    """
    report = {"slides": 0, "pictures": 0, "picture_slides": [], "issues": []}
    try:
        document = PptPresentation(str(pptx_file))
    except Exception as exc:
        report["issues"].append(f"cannot open PPTX: {exc}")
        return report

    width, height = document.slide_width, document.slide_height
    tolerance = 9525 * 2  # 2px in EMU
    report["slides"] = len(document.slides)

    for number, slide in enumerate(document.slides, start=1):
        boxes = []
        for shape in slide.shapes:
            if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.PICTURE:
                report["pictures"] += 1
                report["picture_slides"].append(number)

            if (shape.left < -tolerance or shape.top < -tolerance
                    or shape.left + shape.width > width + tolerance
                    or shape.top + shape.height > height + tolerance):
                report["issues"].append(f"slide {number}: '{shape.name}' extends off the slide")

            if shape.name in _CONTENT_SHAPES:
                boxes.append((shape.name, shape.left, shape.top,
                              shape.left + shape.width, shape.top + shape.height))

            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.font.size is not None and run.font.size.pt < 9:
                            report["issues"].append(
                                f"slide {number}: font {run.font.size.pt}pt in '{shape.name}'"
                            )

        for i, first in enumerate(boxes):
            for second in boxes[i + 1:]:
                overlap_x = min(first[3], second[3]) - max(first[1], second[1])
                overlap_y = min(first[4], second[4]) - max(first[2], second[2])
                if overlap_x > tolerance and overlap_y > tolerance:
                    report["issues"].append(
                        f"slide {number}: '{first[0]}' overlaps '{second[0]}'"
                    )

    return report


# ============================================================
# ENGINES
# ============================================================

# engine id -> label shown in the terminal
PPTX_ENGINES = {
    "compono": "Compono",
    "python-pptx": "python-pptx",
    "both": "Both (Compono layout + python-pptx brand pass)",
}


class PptxGenerationError(RuntimeError):
    """Rendering failed or did not produce the PPTX file."""

    stage = "generation"


class PptxValidationError(RuntimeError):
    """Office-Kit could not validate the final PPTX (the file is kept)."""

    stage = "Office-Kit validation"

    def __init__(self, message: str, review: dict | None = None):
        super().__init__(message)
        self.review = review


def _render_compono(output_file: Path, data: dict) -> Path:
    from office.compono_pptx import render_compono_pptx

    return render_compono_pptx(output_file, data, brand_pass=False)


def _render_both(output_file: Path, data: dict) -> Path:
    # Compono writes the file; python-pptx then reopens that saved file
    # (not Compono's objects) and applies the brand pass in place.
    from office.compono_pptx import render_compono_pptx

    return render_compono_pptx(output_file, data, brand_pass=True)


def _render_python_pptx(output_file: Path, data: dict) -> Path:
    with PERF.stage("PPTX slide building + save (python-pptx)"):
        return render_rich_pptx(output_file=output_file, presentation_data=data)


_RENDERERS = {
    "compono": _render_compono,
    "python-pptx": _render_python_pptx,
    "both": _render_both,
}


def _slide_dicts(data: dict) -> list[dict]:
    return [s if isinstance(s, dict) else s.model_dump() for s in data.get("slides", [])]


def render_pptx(presentation: Any, output_file: Path, engine: str | None = None) -> Path:
    """
    Render the presentation with one engine and return the saved file.

    engine=None uses config.PPTX_ENGINE and, as before, falls back to the
    python-pptx renderer if a Compono-based engine fails. An explicit engine
    never falls back: its failure raises PptxGenerationError.
    """
    use_fallback = engine is None
    engine = engine or PPTX_ENGINE
    if engine not in _RENDERERS:
        raise ValueError(f"Unknown PPTX engine '{engine}'. Choose one of: {', '.join(_RENDERERS)}.")

    output_file = Path(output_file)
    if output_file.suffix.lower() != ".pptx":
        raise PptxGenerationError(f"Output path must end with .pptx: {output_file}")
    try:
        output_file.parent.mkdir(parents=True, exist_ok=True)
    except OSError as exc:
        raise PptxGenerationError(f"Invalid output path {output_file}: {exc}") from exc

    data = _to_dict(presentation)
    if not data.get("slides"):
        raise ValueError(
            "The presentation contains no slides. "
            "Generate slide content before rendering."
        )

    print(f"\n[1/3] Rendering PowerPoint slides (engine: {PPTX_ENGINES[engine]})...")

    try:
        try:
            result = _RENDERERS[engine](output_file, data)
            PERF.incr(f"rendered with {engine}")
        except Exception as exc:
            if not (use_fallback and engine != "python-pptx"):
                raise
            print(f"[warning] Compono rendering failed, using python-pptx renderer: {exc!r}")
            PERF.incr("compono fallbacks")
            result = _render_python_pptx(output_file, data)
    except Exception as exc:
        raise PptxGenerationError(f"PowerPoint rendering failed ({engine}): {exc}") from exc

    if result is not None:
        output_file = Path(result)
    if not output_file.is_file():
        raise PptxGenerationError(
            f"The renderer finished without creating the PPTX file: {output_file}"
        )

    print(f"PPTX created: {output_file}")
    return output_file


def validate_pptx(output_file: Path, presentation: Any) -> dict:
    """
    Office-Kit validation of the final PPTX (open check, stats, issues,
    every slide title present, placeholder repair). Runs once per file.

    Raises PptxValidationError when the file is missing, Office-Kit raises,
    or the review is not valid. The PPTX is never deleted.
    """
    output_file = Path(output_file)
    data = _to_dict(presentation)
    title = data.get("title", "Presentation")

    if not output_file.is_file():
        raise PptxValidationError(f"PPTX file to validate does not exist: {output_file}")

    print("\n[2/3] Inspecting and validating PPTX with Office-Kit...")
    expected_titles = [str(s.get("title") or "") for s in _slide_dicts(data)]
    try:
        with PERF.stage("Office-Kit inspect + validate"):
            review = review_document(
                output_file,
                expected_titles=expected_titles,
                placeholder_values={"title": title},
            )
    except Exception as exc:
        raise PptxValidationError(f"Office-Kit raised {exc!r} while validating {output_file}") from exc

    print_review(review)
    if not review["valid"]:
        problems = review["issues"] + [f"missing title: {t}" for t in review["missing_titles"]]
        raise PptxValidationError(
            f"Office-Kit validation failed for {output_file}: {problems}", review
        )
    return review


def _check_layout_and_images(output_file: Path, data: dict) -> None:
    """Project's own layout/image report (warnings only, not Office-Kit)."""

    print("\n[3/3] Checking layout and embedded images...")
    validation_started = time.perf_counter()

    expected_images = 0
    for slide_dict in _slide_dicts(data):
        image_path = slide_dict.get("image_path")
        if image_path and Path(image_path).is_file():
            expected_images += 1

    layout_report = check_pptx_layout(output_file)
    embedded_images = layout_report["pictures"]
    validation_elapsed = time.perf_counter() - validation_started
    PERF.add_time("PPTX final validation", validation_elapsed)
    print(f"[timing] PPTX final validation: {validation_elapsed:.2f}s")

    PERF.incr("slides generated", layout_report["slides"])
    PERF.incr("images inserted into PPTX", embedded_images)

    print(f"[verification] Slides in PPTX: {layout_report['slides']}")
    print(
        f"[verification] Embedded images in PPTX: {embedded_images} / expected {expected_images}"
        f" (slides {layout_report['picture_slides']})"
    )
    if expected_images and embedded_images < expected_images:
        print("[warning] Some slides were expected to contain images, but fewer pictures were embedded than expected.")
    if layout_report["issues"]:
        print("[warning] Layout check found issues:")
        for issue in layout_report["issues"]:
            print(f"  - {issue}")
    else:
        print("[verification] Layout check: no off-slide shapes, overlaps or tiny fonts.")


def generate_pptx_with_officekit(
    presentation: Any,
    output_file: Path,
    engine: str | None = None,
) -> Path:
    """
    Generate a PPTX with the chosen engine, then validate it with Office-Kit.

    Args:
        presentation: Presentation Pydantic model or dictionary.
        output_file: Destination path for the PPTX file.
        engine: "compono", "python-pptx", "both", or None for
            config.PPTX_ENGINE (with the python-pptx fallback).

    Returns:
        Path to the generated and validated PPTX file.

    Raises:
        TypeError: If the presentation data has an unsupported type.
        ValueError: If required presentation content is missing.
        PptxGenerationError: If rendering fails or creates no file.
        PptxValidationError: If Office-Kit validation fails (file is kept).
    """

    data = _to_dict(presentation)
    print("\n========================================")
    print("       PROFESSIONAL PPTX GENERATOR")
    print("========================================")
    print(f"Presentation: {data.get('title', 'Presentation')}")
    print(f"Number of slides: {len(data.get('slides', []))}")
    print(f"Output file: {output_file}")

    print("\nGenerating presentation...")
    output_file = render_pptx(data, output_file, engine)
    print("PPTX generation completed.")

    print("\nRunning Office-Kit validation...")
    validate_pptx(output_file, data)
    print("Validation passed.")

    _check_layout_and_images(output_file, data)

    print("\n========================================")
    print("       PPTX GENERATION COMPLETED")
    print("========================================")
    print(f"File: {output_file}")

    return output_file


def generate_pptx(input_data: Any, output_path: Path, engine: str | None = None) -> Path:
    """Common interface: generate + validate, return the final PPTX path."""
    return generate_pptx_with_officekit(input_data, output_path, engine)
