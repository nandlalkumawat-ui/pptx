"""
Compono PPTX renderer.

Pipeline:
1. Presentation schema -> one layout plan entry per slide, using the
   existing layout rules (_choose_layout) and references builder.
2. Plan -> Compono spec (header / text / image / grid / sequence / table /
   chart / shape primitives). Compono's resolver computes every position.
3. compono.validate() measures text overflow with real glyph metrics.
   Overflowing slides are split (or trimmed into speaker notes) and
   re-validated, so no slide is rendered with clipped text.
4. Images are pre-cropped to the exact rect Compono resolved for them, so
   nothing is stretched (Compono 0.5.0's fit="cover" scales to the box
   without cropping).
5. compono.render_deck() writes the .pptx.
6. A brand pass (python-pptx) restyles the shapes Compono created and adds
   what Compono 0.5.0 does not draw: credits under real pictures and the
   title-slide hero image.
"""

import copy
import hashlib
import math
import re
from pathlib import Path
from typing import Any

from compono import Deck, render_deck, validate
from compono.resolver import Template, resolve_slide
from compono.schema import Chart, Header, Image, Sequence, Shape, Table, Text
from PIL import Image as PILImage, ImageOps
from pptx import Presentation as PptPresentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from config import BRAND_ACCENT, BRAND_FONT, BRAND_PRIMARY, OUTPUT_DIR
from office.officekit_utils import (
    _LEAD_TERM,
    _choose_layout,
    _references_slide,
    _short_url,
    estimate_text_height,
    fit_font_size,
    to_dict,
)
from utils.perf import PERF


EMU_PER_INCH = 914400

PAGE_W_IN = 13.333
PAGE_H_IN = 7.5
MARGIN_X_IN = 0.6

TITLE_COLOR = BRAND_PRIMARY
SUBTITLE_COLOR = "#475569"
BODY_COLOR = "#1F2937"
MUTED_COLOR = "#64748B"
PANEL_FILL = "#EEF2F7"
ON_DARK_SUBTITLE = "#CBD5E1"

HERO_TITLE_MAX_PT = 40
SUBTITLE_PT = 18
STEP_DESCRIPTION_PT = 16
CARD_LEAD_PT = 20
CARD_BODY_PT = 16
CAPTION_PT = 10
REFERENCE_PT = 12

CAPTION_H_IN = 0.32
HERO_IMAGE_X_IN = 7.0
REFS_PER_SLIDE = 7
MAX_FIT_ROUNDS = 12

IMAGE_ID = "slide_image"
CROP_DIR = OUTPUT_DIR / "images" / "cropped"


class ComponoLayoutError(RuntimeError):
    """The deck cannot be expressed as a valid Compono spec."""


def _emu(inches: float) -> int:
    return round(inches * EMU_PER_INCH)


def build_template() -> Template:
    """16:9 template with the project's brand colours (no site-packages edits)."""
    return Template(
        page_width=_emu(PAGE_W_IN),
        page_height=_emu(PAGE_H_IN),
        margin_top=_emu(0.45),
        margin_right=_emu(MARGIN_X_IN),
        margin_bottom=_emu(0.3),
        margin_left=_emu(MARGIN_X_IN),
        header_height=_emu(1.35),
        footer_height=_emu(0.35),
        gutter=_emu(0.3),
        font_family=BRAND_FONT,
        primary_color=BRAND_PRIMARY,
        accent_color=BRAND_ACCENT,
    )


# ============================================================
# PLAN: schema slides -> layout decisions
# ============================================================

def _notes(slide: dict) -> str:
    parts = []
    if slide.get("notes"):
        parts.append(str(slide["notes"]))
    if slide.get("image_url"):
        parts.append(f"Image source: {slide['image_url']}")
    if slide.get("source_urls"):
        parts.append("Sources:\n" + "\n".join(map(str, slide["source_urls"])))
    return "\n\n".join(parts)


def _plan(presentation_data: dict) -> list[dict]:
    slides = [to_dict(slide) for slide in presentation_data.get("slides", [])]
    plan = []
    for index, slide in enumerate(slides, start=1):
        image_path = str(slide.get("image_path") or "")
        if image_path and not Path(image_path).is_file():
            print(f"[warning] Image file missing for slide {index}, rendering without it: {image_path}")
            image_path = ""
            slide["image_path"] = ""

        caption = str(slide.get("image_caption") or "")
        if image_path and not caption and slide.get("image_source"):
            caption = f"Image: {slide['image_source']}"

        table = slide.get("table")
        plan.append({
            "kind": _choose_layout(slide, index),
            "title": str(slide.get("title") or f"Slide {index}"),
            "subtitle": str(slide.get("subtitle") or ""),
            "paragraph": " ".join(str(p) for p in slide.get("paragraphs") or [] if p),
            "bullets": [str(b) for b in slide.get("bullets") or [] if b],
            "table": {"headers": list(table["headers"]), "rows": [list(r) for r in table["rows"]]}
            if table else None,
            "image_path": image_path,
            "image_caption": caption,
            "notes": _notes(slide),
        })

    references = _references_slide(presentation_data, slides)
    if references:
        rows = [[label, _short_url(url) if url else ""] for label, url in references["_references"]]
        for start in range(0, len(rows), REFS_PER_SLIDE):
            plan.append({
                "kind": "references",
                "title": references["title"] + (" (continued)" if start else ""),
                "subtitle": "" if start else references["subtitle"],
                "paragraph": "",
                "bullets": [],
                "table": {"headers": ["Source", "Link"], "rows": rows[start:start + REFS_PER_SLIDE]},
                "image_path": "",
                "image_caption": "",
                "notes": "\n".join(url for _, url in references["_references"] if url),
            })
    return plan


# ============================================================
# SPEC: plan -> Compono primitives
# ============================================================

def _text(mode: str, content) -> dict:
    return {"primitive": "text", "mode": mode, "content": content}


def _column(items: list[dict]) -> dict:
    if len(items) == 1:
        return items[0]
    return {"primitive": "grid", "columns": 1, "direction": "column", "items": items}


def _row(items: list[dict]) -> dict:
    if len(items) == 1:
        return items[0]
    return {"primitive": "grid", "columns": len(items), "items": items}


def _split_lead(text: str) -> tuple[str, str]:
    match = _LEAD_TERM.match(text)
    if match and len(match.group(1).split()) <= 6:
        return match.group(1).strip(), match.group(2).strip()
    return "", text.strip()


def _card(text: str, fill: str, color: str) -> dict:
    lead, rest = _split_lead(text)
    return {
        "primitive": "shape",
        "kind": "rounded_rect",
        "fill": fill,
        "text": {"content": f"{lead}\n{rest}" if lead else rest,
                 "align": "left", "valign": "top", "color": color},
    }


def _card_grid(bullets: list[str], fill: str, color: str) -> dict:
    count = len(bullets)
    columns = 2 if count in (2, 4) else min(count, 3)
    return {"primitive": "grid", "columns": columns,
            "items": [_card(bullet, fill, color) for bullet in bullets]}


_NUMERIC = re.compile(r"^[^\d\-+]{0,3}([-+]?\d[\d,]*(?:\.\d+)?)\s*(%|[A-Za-z]{0,5})?$")


def _number(cell: str) -> float | None:
    match = _NUMERIC.match(str(cell).strip())
    if not match:
        return None
    try:
        return float(match.group(1).replace(",", ""))
    except ValueError:
        return None


def _chart_for(table: dict) -> dict | None:
    """A native bar chart when the table is category + numeric columns."""
    rows = table["rows"]
    if not 2 <= len(rows) <= 10 or len(table["headers"]) < 2:
        return None
    series = []
    for col in range(1, min(len(table["headers"]), 4)):
        values = [_number(row[col]) for row in rows]
        if all(value is not None for value in values):
            series.append({"name": table["headers"][col], "values": values})
    if not series:
        return None
    return {"primitive": "chart", "chart_type": "bar",
            "categories": [str(row[0]) for row in rows], "series": series}


def _body(item: dict) -> list[dict]:
    kind = item["kind"]
    paragraph, bullets = item["paragraph"], item["bullets"]
    text_items = []
    if paragraph:
        text_items.append(_text("paragraph", paragraph))
    if bullets:
        text_items.append(_text("bullets", bullets))

    if kind == "title":
        return []

    if kind == "image_text" and item.get("image_src"):
        image = {"primitive": "image", "id": IMAGE_ID, "src": item["image_src"], "fit": "cover"}
        return [_row([_column(text_items), image])] if text_items else [image]

    if kind == "process" and len(bullets) >= 2:
        steps = []
        for number, bullet in enumerate(bullets, start=1):
            lead, rest = _split_lead(bullet)
            steps.append({"label": lead or f"Step {number}", "description": rest})
        sequence = {"primitive": "sequence", "orientation": "horizontal", "steps": steps}
        return ([_text("paragraph", paragraph)] if paragraph else []) + [sequence]

    if kind in ("comparison", "conclusion") and len(bullets) >= 2:
        if kind == "conclusion":
            cards = _card_grid(bullets, BRAND_PRIMARY, "#FFFFFF")
        else:
            cards = _card_grid(bullets, PANEL_FILL, BODY_COLOR)
        return ([_text("paragraph", paragraph)] if paragraph else []) + [cards]

    if kind in ("data", "references") and item.get("table"):
        table = {"primitive": "table", "headers": item["table"]["headers"],
                 "rows": item["table"]["rows"]}
        if kind == "references":
            return [table]
        chart = _chart_for(item["table"])
        left = _column([_text("paragraph", paragraph), table]) if paragraph and chart else table
        right = chart or (_text("bullets", bullets) if bullets else None)
        body = [_row([left, right])] if right else [left]
        if paragraph and not chart:
            body.insert(0, _text("paragraph", paragraph))
        return body

    # Plain text slide: lead paragraph in a panel next to the bullets.
    if paragraph and bullets:
        panel = {"primitive": "shape", "kind": "rounded_rect", "fill": PANEL_FILL,
                 "text": {"content": paragraph, "align": "left", "valign": "top",
                          "color": BODY_COLOR}}
        return [_row([panel, _text("bullets", bullets)])]
    return text_items


def _header(item: dict) -> dict:
    header = {"title": item["title"]}
    if item["subtitle"]:
        header["subtitle"] = item["subtitle"]
    return header


def _spec(plan: list[dict]) -> dict:
    slides = []
    for item in plan:
        slide = {"header": _header(item), "body": _body(item)}
        if item["notes"]:
            slide["notes"] = item["notes"]
        slides.append(slide)
    return {"slides": slides}


# ============================================================
# FIT: split slides that Compono reports as overflowing
# ============================================================

def _continuation(item: dict, **changes) -> dict:
    follow = copy.deepcopy(item)
    follow.update(title=f"{item['title'].removesuffix(' (continued)')} (continued)",
                  subtitle="", paragraph="", image_path="", image_src="",
                  image_caption="", notes="")
    follow.update(changes)
    return follow


def _kind_for(kind: str, bullets: list[str], has_image: bool) -> str:
    if kind in ("process", "comparison", "conclusion") and len(bullets) >= 2:
        return kind
    if kind in ("process", "comparison", "conclusion", "image_text"):
        return "image_text" if has_image else "text"
    return kind


def _split(item: dict, errors: list[dict]) -> list[dict]:
    """Make one overflowing slide fit, preferring a continuation slide."""
    if any(str(e.get("primitive", "")).startswith("header") for e in errors):
        if item["subtitle"]:
            item["notes"] = f"{item['subtitle']}\n\n{item['notes']}".strip()
            item["subtitle"] = ""
            return [item]
        words = item["title"].split()
        if len(words) > 4:
            item["title"] = " ".join(words[: max(4, len(words) * 2 // 3)])
            return [item]

    table = item.get("table")
    if table and len(table["rows"]) > 2:
        half = math.ceil(len(table["rows"]) / 2)
        first = dict(item, table=dict(table, rows=table["rows"][:half]))
        second = _continuation(item, table=dict(table, rows=table["rows"][half:]), bullets=[])
        return [first, second]

    bullets = item["bullets"]
    if len(bullets) >= 2:
        half = math.ceil(len(bullets) / 2)
        has_image = bool(item.get("image_src"))
        first = dict(item, bullets=bullets[:half],
                     kind=_kind_for(item["kind"], bullets[:half], has_image))
        second = _continuation(item, bullets=bullets[half:],
                               kind=_kind_for(item["kind"], bullets[half:], False))
        return [first, second]

    if item["paragraph"] and bullets:
        first = dict(item, bullets=[], kind=_kind_for(item["kind"], [], bool(item.get("image_src"))))
        second = _continuation(item, bullets=bullets, kind="text")
        return [first, second]

    sentences = re.split(r"(?<=[.!?])\s+", item["paragraph"]) if item["paragraph"] else []
    if len(sentences) > 1:
        half = math.ceil(len(sentences) / 2)
        first = dict(item, paragraph=" ".join(sentences[:half]))
        second = _continuation(item, paragraph=" ".join(sentences[half:]), kind="text")
        return [first, second]

    if item.get("image_src"):
        item.update(image_src="", kind="text")
        return [item]

    raise ComponoLayoutError(f"Slide '{item['title']}' does not fit: {errors}")


def _fit(plan: list[dict], template: Template) -> tuple[list[dict], dict]:
    for _ in range(MAX_FIT_ROUNDS):
        spec = _spec(plan)
        report = validate(spec, template=template)
        if report.valid:
            return plan, spec

        hard = [e for e in report.errors if e.get("error") != "overflow"]
        if hard:
            raise ComponoLayoutError(f"Compono rejected the spec: {hard[:3]}")

        by_slide: dict[int, list[dict]] = {}
        for error in report.errors:
            by_slide.setdefault(int(error["slide"]), []).append(error)
        PERF.incr("compono overflow fixes", len(by_slide))
        for index in sorted(by_slide, reverse=True):
            plan[index:index + 1] = _split(plan[index], by_slide[index])

    raise ComponoLayoutError(f"Text still overflows after {MAX_FIT_ROUNDS} rounds.")


# ============================================================
# IMAGES: crop to the resolved box so nothing is stretched
# ============================================================

def crop_to_ratio(src: str, width: int, height: int) -> str:
    """Centre-crop an image to width:height (cached, never upscaled)."""
    ratio = width / height
    digest = hashlib.sha1(f"{src}|{ratio:.4f}".encode("utf-8")).hexdigest()
    target = CROP_DIR / f"{digest}.jpg"
    if target.is_file() and target.stat().st_size > 0:
        PERF.incr("cropped image cache hits")
        return str(target)

    CROP_DIR.mkdir(parents=True, exist_ok=True)
    with PILImage.open(src) as image:
        image = image.convert("RGB")
        if image.width / image.height > ratio:
            size = (round(image.height * ratio), image.height)
        else:
            size = (image.width, round(image.width / ratio))
        cropped = ImageOps.fit(image, size, method=PILImage.LANCZOS, centering=(0.5, 0.4))
        temp = target.with_suffix(".tmp")
        cropped.save(temp, format="JPEG", quality=88)
        temp.replace(target)
    return str(target)


def _prepare_images(plan: list[dict], template: Template) -> None:
    """Pre-crop every content image to the box Compono will give it."""
    for item in plan:
        if item["kind"] == "image_text" and item["image_path"]:
            item["image_src"] = item["image_path"]  # placeholder for resolving

    deck = Deck.model_validate(_spec(plan))
    for item, slide in zip(plan, deck.slides):
        if not item.get("image_src"):
            continue
        layout = resolve_slide(template, header=slide.header, body=slide.body)
        rect = layout.rects[IMAGE_ID]
        caption_h = _emu(CAPTION_H_IN) if item["image_caption"] else 0
        item["image_src"] = crop_to_ratio(item["image_path"], rect.w, rect.h - caption_h)


# ============================================================
# BRAND PASS (python-pptx on the file Compono wrote)
# ============================================================

def _rgb(color: str) -> RGBColor:
    return RGBColor.from_string(color.lstrip("#"))


def _font(font, size: float | None = None, color: str | None = None,
          bold: bool | None = None) -> None:
    font.name = BRAND_FONT
    if size is not None:
        font.size = Pt(size)
    if color is not None:
        font.color.rgb = _rgb(color)
    if bold is not None:
        font.bold = bold


def _shapes_at(slide, rect) -> list:
    return [shape for shape in slide.shapes
            if (shape.left, shape.top, shape.width, shape.height) == (rect.x, rect.y, rect.w, rect.h)]


def _restyle_text_frame(frame, color: str, lead_color: str, anchor_top: bool = True) -> None:
    """Body colour, top anchoring and bold 'Term:' leads."""
    if anchor_top:
        frame.vertical_anchor = MSO_ANCHOR.TOP
    for paragraph in frame.paragraphs:
        text = paragraph.text
        bullet = text.startswith("• ")
        lead, rest = _split_lead(text[2:] if bullet else text)
        size = paragraph.runs[0].font.size if paragraph.runs else None
        paragraph.space_after = Pt(8)
        if lead:
            for run in list(paragraph.runs):
                run._r.getparent().remove(run._r)
            first = paragraph.add_run()
            first.text = ("• " if bullet else "") + lead + ": "
            _font(first.font, color=lead_color, bold=True)
            second = paragraph.add_run()
            second.text = rest
            _font(second.font, color=color, bold=False)
            for run in (first, second):
                if size is not None:
                    run.font.size = size
        else:
            for run in paragraph.runs:
                _font(run.font, color=color)


def _style_header(shape, dark: bool = False) -> None:
    shape.name = "Slide Title"
    paragraphs = shape.text_frame.paragraphs
    _font(paragraphs[0].font, color="#FFFFFF" if dark else TITLE_COLOR, bold=True)
    for paragraph in paragraphs[1:]:
        _font(paragraph.font, size=SUBTITLE_PT, color=ON_DARK_SUBTITLE if dark else SUBTITLE_COLOR)


def _add_rect(slide, x: int, y: int, w: int, h: int, color: str, name: str):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(h))
    bar.fill.solid()
    bar.fill.fore_color.rgb = _rgb(color)
    bar.line.fill.background()
    bar.name = name
    return bar


def _add_caption(slide, text: str, x: int, y: int, w: int, h: int, color: str = MUTED_COLOR,
                 name: str = "Image Caption") -> None:
    box = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    box.name = name
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.margin_left = frame.margin_right = 0
    frame.text = text
    _font(frame.paragraphs[0].font, size=CAPTION_PT, color=color)
    frame.paragraphs[0].alignment = PP_ALIGN.RIGHT


def _style_title_slide(slide, item: dict, header_shape, template: Template, deck_title: str) -> None:
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb(BRAND_PRIMARY)
    _style_header(header_shape, dark=True)
    header_shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE

    width_in = PAGE_W_IN - 2 * MARGIN_X_IN
    if item["image_path"]:
        width_in = HERO_IMAGE_X_IN - MARGIN_X_IN - 0.4
        header_shape.width = Emu(_emu(width_in))
        hero_w, hero_h = template.page_width - _emu(HERO_IMAGE_X_IN), template.page_height
        hero = crop_to_ratio(item["image_path"], hero_w, hero_h)
        picture = slide.shapes.add_picture(hero, Emu(_emu(HERO_IMAGE_X_IN)), 0, Emu(hero_w), Emu(hero_h))
        picture.name = "Slide Image"
        if item["image_caption"]:
            _add_caption(slide, item["image_caption"], _emu(HERO_IMAGE_X_IN) + _emu(0.2),
                         hero_h - _emu(0.45), hero_w - _emu(0.4), _emu(0.35), "#FFFFFF",
                         name="Hero Credit")

    title_size = fit_font_size([item["title"]], width_in, 2.6, max_size=HERO_TITLE_MAX_PT, min_size=28) or 28
    _font(header_shape.text_frame.paragraphs[0].font, size=title_size)
    _add_rect(slide, _emu(MARGIN_X_IN), _emu(PAGE_H_IN / 2) - _emu(1.9), _emu(1.2), _emu(0.08),
              BRAND_ACCENT, "Accent Bar")


def _style_footer(slide, template: Template, deck_title: str, dark: bool) -> None:
    footer_top = template.page_height - template.margin_bottom - template.footer_height
    for shape in slide.shapes:
        if shape.top == footer_top and shape.has_text_frame and "/" in shape.text_frame.text:
            shape.name = "Footer"
            _font(shape.text_frame.paragraphs[0].font, color=MUTED_COLOR)
    if deck_title:
        box = slide.shapes.add_textbox(Emu(template.margin_left), Emu(footer_top),
                                       Emu(_emu(8.0)), Emu(template.footer_height))
        box.name = "Footer Title"
        frame = box.text_frame
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        frame.margin_left = 0
        frame.text = deck_title
        _font(frame.paragraphs[0].font, size=10, color=MUTED_COLOR)


def _style_sequence(slide, rect) -> None:
    """Make step boxes a sensible height and step text readable."""
    max_h = _emu(3.0)
    for shape in slide.shapes:
        inside = (shape.left >= rect.x and shape.top >= rect.y
                  and shape.left + shape.width <= rect.x + rect.w
                  and shape.top + shape.height <= rect.y + rect.h)
        if not inside or shape.shape_type != MSO_SHAPE_TYPE.AUTO_SHAPE or not shape.has_text_frame:
            continue
        shape.name = "Process Step"
        if shape.height > max_h:
            shape.top = shape.top + (shape.height - max_h) // 2
            shape.height = max_h
        paragraphs = shape.text_frame.paragraphs
        _font(paragraphs[0].font, color="#FFFFFF", bold=True)
        width_in = shape.width / EMU_PER_INCH
        description = [p.text for p in paragraphs[1:]]
        size = fit_font_size(description, width_in, shape.height / EMU_PER_INCH - 0.6,
                             max_size=STEP_DESCRIPTION_PT, min_size=12) or 12
        for paragraph in paragraphs[1:]:
            _font(paragraph.font, size=size, color="#FFFFFF")


def _style_card(shape) -> None:
    shape.name = "Card"
    frame = shape.text_frame
    frame.margin_left = frame.margin_right = Inches(0.2)
    frame.margin_top = Inches(0.18)
    paragraphs = frame.paragraphs
    multi = len(paragraphs) > 1
    if not multi:
        # Single-statement panel (lead paragraph): larger, vertically centred.
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        frame.margin_left = frame.margin_right = Inches(0.35)
        size = fit_font_size([frame.text], shape.width / EMU_PER_INCH - 0.7,
                             shape.height / EMU_PER_INCH - 0.5, max_size=22,
                             min_size=CARD_BODY_PT) or CARD_BODY_PT
        _font(paragraphs[0].font, size=size)
        return
    for number, paragraph in enumerate(paragraphs):
        lead = multi and number == 0
        _font(paragraph.font, size=CARD_LEAD_PT if lead else CARD_BODY_PT, bold=lead)
        paragraph.space_after = Pt(6)


def _tighten_below_paragraph(slide, paragraph_shape, footer_top: int) -> None:
    """
    Compono 0.5.0 splits stacked items equally, so a short paragraph above
    bullets/cards/tables leaves a large gap. Shrink the paragraph box to its
    text and move what sits below it (same column) up into the freed space.
    """
    width_in = paragraph_shape.width / EMU_PER_INCH
    needed = _emu(estimate_text_height([paragraph_shape.text_frame.text], width_in, 18) + 0.1)
    old_bottom = paragraph_shape.top + paragraph_shape.height
    delta = paragraph_shape.height - needed
    if delta <= _emu(0.3):
        return
    paragraph_shape.height = Emu(needed)

    left, right = paragraph_shape.left, paragraph_shape.left + paragraph_shape.width
    for shape in slide.shapes:
        if shape is paragraph_shape or shape.top < old_bottom - _emu(0.05) or shape.top >= footer_top:
            continue
        if shape.left + shape.width <= left or shape.left >= right:
            continue  # another column
        if shape.name in ("Bullet List", "Card"):
            shape.top = shape.top - delta
            shape.height = shape.height + delta
        elif shape.name == "Content Table" or shape.name == "Chart":
            shape.top = shape.top - delta
        else:  # process steps and their connectors: recentre in the larger area
            shape.top = shape.top - delta // 2


def _style_table(frame, references: bool) -> None:
    frame.name = "Content Table"
    table = frame.table
    row_h = _emu(0.42)
    for row in table.rows:
        row.height = Emu(row_h)
    frame.height = Emu(row_h * len(table.rows))
    if references:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.text_frame.paragraphs:
                    _font(paragraph.font, size=REFERENCE_PT)


def _style_chart(frame) -> None:
    frame.name = "Chart"
    chart = frame.chart
    chart.font.size = Pt(14)
    chart.font.name = BRAND_FONT
    palette = [BRAND_PRIMARY, BRAND_ACCENT, "#94A3B8"]
    for number, series in enumerate(chart.plots[0].series):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = _rgb(palette[number % len(palette)])


def _brand_pass(pptx_file: Path, plan: list[dict], render_report, template: Template,
                deck_title: str) -> None:
    prs = PptPresentation(str(pptx_file))
    for slide, item, layout in zip(prs.slides, plan, render_report.actual_layout):
        title_slide = item["kind"] == "title"
        for item_id, rect in layout.rects.items():
            primitive = layout.items.get(item_id)
            shapes = _shapes_at(slide, rect)

            if isinstance(primitive, Header):
                for shape in shapes:
                    if title_slide:
                        _style_title_slide(slide, item, shape, template, deck_title)
                    else:
                        _style_header(shape)
                        _add_rect(slide, _emu(0.36), rect.y + _emu(0.3), _emu(0.07),
                                  rect.h - _emu(0.6), BRAND_ACCENT, "Accent Bar")
            elif isinstance(primitive, Text):
                for shape in shapes:
                    shape.name = "Bullet List" if primitive.mode == "bullets" else "Body Paragraph"
                    # Next to a centred lead panel, centre the bullets too.
                    beside_panel = item["kind"] == "text" and item["paragraph"] and primitive.mode == "bullets"
                    _restyle_text_frame(shape.text_frame, BODY_COLOR, TITLE_COLOR,
                                        anchor_top=not beside_panel)
                    if beside_panel:
                        shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            elif isinstance(primitive, Shape):
                for shape in shapes:
                    _style_card(shape)
            elif isinstance(primitive, Table):
                for shape in shapes:
                    _style_table(shape, item["kind"] == "references")
            elif isinstance(primitive, Chart):
                for shape in shapes:
                    _style_chart(shape)
            elif isinstance(primitive, Sequence):
                _style_sequence(slide, rect)
            elif isinstance(primitive, Image):
                for shape in shapes:
                    shape.name = "Slide Image"
                    if item["image_caption"]:
                        shape.height = Emu(rect.h - _emu(CAPTION_H_IN))
                        _add_caption(slide, item["image_caption"], rect.x,
                                     rect.y + rect.h - _emu(CAPTION_H_IN), rect.w, _emu(CAPTION_H_IN))

        footer_top = template.page_height - template.margin_bottom - template.footer_height
        for shape in list(slide.shapes):
            if shape.name == "Body Paragraph":
                _tighten_below_paragraph(slide, shape, footer_top)

        if title_slide:
            for shape in list(slide.shapes):
                if shape.has_text_frame and re.fullmatch(r"\d+ / \d+", shape.text_frame.text.strip()):
                    shape._element.getparent().remove(shape._element)
        else:
            _style_footer(slide, template, deck_title, dark=False)

    prs.save(str(pptx_file))


# ============================================================
# PUBLIC ENTRY POINT
# ============================================================

def render_compono_pptx(output_file: Path, presentation_data: Any) -> Path:
    """Render a Presentation (model or dict) to .pptx with Compono."""
    output_file = Path(output_file)
    output_file.parent.mkdir(parents=True, exist_ok=True)

    data = to_dict(presentation_data)
    if not data.get("slides"):
        raise ValueError("Presentation contains no slides.")

    template = build_template()
    deck_title = str(data.get("title") or "")

    with PERF.stage("Compono layout plan + validate"):
        plan = _plan(data)
        _prepare_images(plan, template)
        plan, spec = _fit(plan, template)
        PERF.incr("compono slides", len(plan))

    with PERF.stage("Compono render_deck (write .pptx)"):
        report = render_deck(spec, str(output_file), template=template)
    if report.warnings:
        print(f"[compono] warnings: {report.warnings}")

    with PERF.stage("Brand pass (python-pptx)"):
        _brand_pass(output_file, plan, report, template, deck_title)

    print(f"PPTX saved successfully (Compono): {output_file}")
    return output_file
