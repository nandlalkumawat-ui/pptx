"""
DOCX report renderer (python-docx).

Report schema -> Word document with:
- core metadata (title, subject, author, keywords)
- A4/Letter page size and margins from config
- consistent styles: Normal, Title, Subtitle, Heading 1-3, Caption, lists
- title page, then a Table of Contents field (filled when Word/LibreOffice
  updates fields on open)
- header (report title) and footer "Page X of Y" via PAGE/NUMPAGES fields,
  suppressed on the title page
- paragraphs, bulleted lists, formatted tables with repeating header rows
- figures with numbered captions (SEQ fields)
- numbered references with real external hyperlinks

PPTX layout logic is deliberately not used here.
"""

import datetime as dt
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor

from config import (
    BRAND_ACCENT,
    BRAND_FONT,
    BRAND_PRIMARY,
    DOCX_BODY_FONT_SIZE,
    DOCX_MARGIN_IN,
    DOCX_PAGE_SIZE,
)


BODY_COLOR = "1F2937"
MUTED_COLOR = "64748B"
TABLE_STRIPE = "EEF2F7"


def _hex(color: str) -> RGBColor:
    return RGBColor.from_string(color.lstrip("#"))


def _to_dict(value: Any) -> dict:
    if hasattr(value, "model_dump"):
        return value.model_dump()
    if isinstance(value, dict):
        return value
    raise TypeError("Report must be a Pydantic model or dictionary.")


# ============================================================
# OOXML helpers (fields, shading, hyperlinks)
# ============================================================

def _add_field(paragraph, instruction: str, placeholder: str = "") -> None:
    """Insert a complex field (PAGE, NUMPAGES, TOC, SEQ ...) into a paragraph."""

    def char(kind: str):
        run = OxmlElement("w:r")
        fld = OxmlElement("w:fldChar")
        fld.set(qn("w:fldCharType"), kind)
        if kind == "begin":
            fld.set(qn("w:dirty"), "true")
        run.append(fld)
        return run

    paragraph._p.append(char("begin"))
    instr_run = OxmlElement("w:r")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    instr_run.append(instr)
    paragraph._p.append(instr_run)
    paragraph._p.append(char("separate"))
    result = paragraph.add_run(placeholder)
    paragraph._p.append(result._r)
    paragraph._p.append(char("end"))


def _shade(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), color.lstrip("#"))
    tc_pr.append(shading)


def _repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_hyperlink(paragraph, url: str, text: str) -> None:
    """External hyperlink run (python-docx has no public API for this)."""
    r_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    style = OxmlElement("w:rStyle")
    style.set(qn("w:val"), "Hyperlink")
    props.append(style)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BRAND_ACCENT.lstrip("#"))
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    text_el = OxmlElement("w:t")
    text_el.text = text
    text_el.set(qn("xml:space"), "preserve")
    run.append(text_el)
    link.append(run)
    paragraph._p.append(link)


def _update_fields_on_open(document) -> None:
    """Ask Word/LibreOffice to refresh TOC/page fields when the file opens."""
    settings = document.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


# ============================================================
# Document setup
# ============================================================

def _setup_styles(document) -> None:
    styles = document.styles

    normal = styles["Normal"]
    normal.font.name = BRAND_FONT
    normal.font.size = Pt(DOCX_BODY_FONT_SIZE)
    normal.font.color.rgb = _hex(BODY_COLOR)
    # East-Asian font slot too, otherwise Word may substitute.
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BRAND_FONT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, color, before, after in (
        ("Title", 30, BRAND_PRIMARY, 0, 8),
        ("Subtitle", 15, MUTED_COLOR, 0, 12),
        ("Heading 1", 18, BRAND_PRIMARY, 18, 6),
        ("Heading 2", 14, BRAND_PRIMARY, 12, 4),
        ("Heading 3", 12, BRAND_ACCENT, 10, 3),
    ):
        style = styles[name]
        style.font.name = BRAND_FONT
        style.font.size = Pt(size)
        style.font.color.rgb = _hex(color)
        style.font.bold = name != "Subtitle"
        style.font.italic = False
        rfonts = style.element.rPr.rFonts
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rfonts.attrib.pop(qn(attr), None)
        rfonts.set(qn("w:ascii"), BRAND_FONT)
        rfonts.set(qn("w:hAnsi"), BRAND_FONT)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        if name.startswith("Heading"):
            style.paragraph_format.keep_with_next = True

    caption = styles["Caption"]
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.bold = False
    caption.font.color.rgb = _hex(MUTED_COLOR)


def _setup_page(section) -> None:
    if DOCX_PAGE_SIZE.lower() == "letter":
        section.page_width, section.page_height = Inches(8.5), Inches(11)
    else:
        section.page_width, section.page_height = Mm(210), Mm(297)
    margin = Inches(DOCX_MARGIN_IN)
    section.left_margin = section.right_margin = margin
    section.top_margin = section.bottom_margin = margin
    section.header_distance = section.footer_distance = Inches(0.45)


def _setup_header_footer(section, title: str) -> None:
    section.different_first_page_header_footer = True  # clean title page

    header = section.header.paragraphs[0]
    header.text = title
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = _hex(MUTED_COLOR)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Page ")
    _add_field(footer, "PAGE", "1")
    footer.add_run(" of ")
    _add_field(footer, "NUMPAGES", "1")
    for run in footer.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = _hex(MUTED_COLOR)


# ============================================================
# Content blocks
# ============================================================

def _bullets(document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        lead, sep, rest = str(item).partition(": ")
        if sep and len(lead.split()) <= 6:
            paragraph.add_run(lead + ": ").bold = True
            paragraph.add_run(rest)
        else:
            paragraph.add_run(str(item))
        paragraph.paragraph_format.space_after = Pt(3)


def _caption(document, label: str, text: str) -> None:
    paragraph = document.add_paragraph(style="Caption")
    paragraph.add_run(f"{label} ")
    number = document._figure_counters.get(label, 0) + 1
    document._figure_counters[label] = number
    _add_field(paragraph, f"SEQ {label} \\* ARABIC", str(number))
    paragraph.add_run(f": {text}")


def _table(document, table_data: dict, caption: str) -> None:
    headers = table_data["headers"]
    rows = table_data["rows"]
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    header_row = table.rows[0]
    _repeat_header(header_row)
    for cell, text in zip(header_row.cells, headers):
        cell.text = ""
        run = cell.paragraphs[0].add_run(str(text))
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _shade(cell, BRAND_PRIMARY)

    for index, row in enumerate(rows):
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(text))
            run.font.size = Pt(10)
            if index % 2:
                _shade(cell, TABLE_STRIPE)

    _caption(document, "Table", caption or "Key figures")


def _figure(document, section: dict, width_in: float) -> bool:
    path = section.get("image_path")
    if not path or not Path(path).is_file():
        return False
    document.add_picture(str(path), width=Inches(width_in))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.paragraphs[-1].paragraph_format.keep_with_next = True
    credit = section.get("image_source") or ""
    text = section.get("image_caption") or section.get("title", "")
    if credit and credit not in text:
        text = f"{text} (Image: {credit})"
    _caption(document, "Figure", text)
    return True


def _references(document, report: dict) -> None:
    document.add_heading("References", level=1)
    entries = []
    for source in report.get("sources") or []:
        source = _to_dict(source)
        if source.get("url") or source.get("title"):
            entries.append((str(source.get("title") or source.get("url")), str(source.get("url") or "")))
    for section in report.get("sections") or []:
        if section.get("image_path") and section.get("image_url"):
            page = next(iter(section.get("source_urls") or []), "") or section["image_url"]
            entries.append((f"Figure image: {section.get('image_source') or 'web'} ({section['title']})", page))

    if not entries:
        document.add_paragraph("No external sources were used.")
        return
    for title, url in entries:
        paragraph = document.add_paragraph(style="List Number")
        paragraph.add_run(title + ". ")
        if url:
            add_hyperlink(paragraph, url, url)
        paragraph.paragraph_format.space_after = Pt(3)


# ============================================================
# PUBLIC ENTRY POINT
# ============================================================

def render_docx_report(output_file: Path, report_data: Any) -> Path:
    """Render a Report (model or dict) to a .docx file."""

    output_file = Path(output_file)
    output_file.parent.mkdir(parents=True, exist_ok=True)
    report = _to_dict(report_data)
    report["sections"] = [_to_dict(s) for s in report.get("sections") or []]
    if not report["sections"]:
        raise ValueError("Report contains no sections.")

    title = str(report.get("title") or "Report")
    document = Document()
    document._figure_counters = {}

    props = document.core_properties
    props.title = title
    props.subject = str(report.get("subtitle") or "")
    props.author = "AI Document Generator"
    props.keywords = ", ".join(s["title"] for s in report["sections"][:6])
    props.created = dt.datetime.now()

    _setup_styles(document)
    section = document.sections[0]
    _setup_page(section)
    _setup_header_footer(section, title)
    text_width = (section.page_width - section.left_margin - section.right_margin) / 914400

    # Title page
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_before = Pt(160)
    document.add_paragraph(title, style="Title")
    if report.get("subtitle"):
        document.add_paragraph(str(report["subtitle"]), style="Subtitle")
    date = document.add_paragraph(dt.date.today().strftime("%d %B %Y"))
    date.runs[0].font.color.rgb = _hex(MUTED_COLOR)
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # Table of contents
    document.add_paragraph("Contents", style="Heading 1")
    toc = document.add_paragraph()
    _add_field(toc, 'TOC \\o "1-3" \\h \\z \\u',
               "Right-click and choose Update Field to build the table of contents.")
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # Body
    if report.get("executive_summary"):
        document.add_heading("Executive Summary", level=1)
        document.add_paragraph(str(report["executive_summary"]))

    for number, part in enumerate(report["sections"], start=1):
        document.add_heading(f"{number}. {part['title']}", level=1)
        for paragraph in part.get("paragraphs") or []:
            document.add_paragraph(str(paragraph))
        _figure(document, part, min(5.5, text_width))
        if part.get("bullets"):
            document.add_heading("Key points", level=3)
            _bullets(document, part["bullets"])
        for sub_number, sub in enumerate(part.get("subsections") or [], start=1):
            sub = _to_dict(sub)
            document.add_heading(f"{number}.{sub_number} {sub['title']}", level=2)
            for paragraph in sub.get("paragraphs") or []:
                document.add_paragraph(str(paragraph))
            if sub.get("bullets"):
                _bullets(document, sub["bullets"])
        if part.get("table"):
            _table(document, _to_dict(part["table"]), part.get("table_caption", ""))

    if report.get("conclusion"):
        document.add_heading("Conclusion", level=1)
        document.add_paragraph(str(report["conclusion"]))

    references = document.add_section(WD_SECTION.NEW_PAGE)  # references on a new page
    references.different_first_page_header_footer = False  # keep header + page numbers
    _references(document, report)

    _update_fields_on_open(document)
    document.save(str(output_file))
    print(f"DOCX saved successfully: {output_file}")
    return output_file
