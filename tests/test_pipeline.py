"""
Offline integration tests (no Ollama, no network).

Run:  .venv/bin/python -m unittest discover -s tests -v
"""

import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document
from PIL import Image
from pptx import Presentation as PptPresentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

import main
from generators import officekit_pptx_generator
from generators.officekit_docx_generator import generate_docx_with_officekit
from generators.officekit_pptx_generator import check_pptx_layout, generate_pptx_with_officekit
from generators.pptx_generator import PPTXContentGenerator
from office.officekit_checks import review_document
from research import image_search
from research.web_search import compact_research, format_search_results
from schemas.document_schema import Presentation, Report


def make_image(path: Path, size=(1600, 1000), color=(30, 90, 140)) -> str:
    Image.new("RGB", size, color).save(path, "JPEG")
    return str(path)


def slide(number, title, layout="content", **extra) -> dict:
    data = {"slide_number": number, "title": title, "subtitle": "Subtitle text",
            "layout": layout, "paragraphs": [], "bullets": [], "table": None}
    data.update(extra)
    return data


SOURCES = [{"title": "Example source", "url": "https://example.org/a", "snippet": "s"}]


class PptxTests(unittest.TestCase):

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.img = make_image(self.tmp / "photo.jpg")

    def deck(self, slides) -> Presentation:
        return Presentation.model_validate({"title": "Energy Storage", "subtitle": "Test deck",
                                            "slides": slides, "sources": SOURCES})

    def standard_slides(self):
        return [
            slide(1, "Energy Storage", "title", image_path=self.img, image_source="example.org"),
            slide(2, "Why storage matters", paragraphs=["Grids need flexibility as solar grows."],
                  bullets=["Peak shaving: lower demand charges", "Backup: critical loads stay on",
                           "Arbitrage: buy low, sell high"],
                  image_path=self.img, image_source="example.org"),
            slide(3, "Building a project", "process",
                  bullets=["Study: grid connection", "Design: size MW and MWh",
                           "Build: racks and inverters", "Operate: dispatch"]),
            slide(4, "Chemistries compared", "comparison",
                  bullets=["Lithium-ion: dense and cheap", "Flow: long duration",
                           "Sodium-ion: no lithium needed"]),
            slide(5, "Installed capacity", "data", paragraphs=["Capacity is rising fast."],
                  table={"headers": ["Year", "GW"], "rows": [["2021", "16"], ["2022", "28"], ["2023", "45"]]}),
            slide(6, "Conclusion", "conclusion",
                  bullets=["Storage is now core grid infrastructure", "Costs keep falling",
                           "Policy decides the pace"]),
        ]

    def test_compono_deck_layouts_images_and_validation(self):
        out = self.tmp / "deck.pptx"
        generate_pptx_with_officekit(self.deck(self.standard_slides()), out)
        prs = PptPresentation(str(out))
        self.assertAlmostEqual(prs.slide_width / prs.slide_height, 16 / 9, places=3)
        names = [{shape.name for shape in s.shapes} for s in prs.slides]
        self.assertIn("Process Step", names[2])
        self.assertIn("Card", names[3])
        self.assertIn("Chart", names[4])
        self.assertIn("Content Table", names[4])
        self.assertTrue(any("Content Table" in n for n in names[6:]))  # references slide
        report = check_pptx_layout(out)
        self.assertEqual(report["issues"], [])
        self.assertEqual(report["picture_slides"], [1, 2])  # embedded, not linked
        review = review_document(out, expected_titles=[s["title"] for s in self.standard_slides()])
        self.assertTrue(review["valid"], review)

    def test_title_and_body_font_sizes(self):
        out = self.tmp / "fonts.pptx"
        generate_pptx_with_officekit(self.deck(self.standard_slides()), out)
        prs = PptPresentation(str(out))
        for number, s in enumerate(prs.slides, start=1):
            for shape in s.shapes:
                if shape.name == "Slide Title":
                    size = shape.text_frame.paragraphs[0].font.size.pt
                    self.assertTrue(28 <= size <= 40, (number, size))
                    self.assertTrue(shape.text_frame.paragraphs[0].font.bold)

    def test_long_content_is_split_not_clipped(self):
        long_bullets = [f"Point {i}: " + "a fairly long explanation of this point " * 3 for i in range(8)]
        slides = [slide(1, "Deck", "title"),
                  slide(2, "Dense slide", paragraphs=["Long intro sentence one. " * 6 + "End."],
                        bullets=long_bullets)]
        out = self.tmp / "long.pptx"
        generate_pptx_with_officekit(self.deck(slides), out)
        prs = PptPresentation(str(out))
        titles = [shape.text_frame.paragraphs[0].text for s in prs.slides
                  for shape in s.shapes if shape.name == "Slide Title"]
        self.assertTrue(any("(continued)" in t for t in titles), titles)
        text = "\n".join(shape.text_frame.text for s in prs.slides for shape in s.shapes
                         if shape.has_text_frame)
        for bullet in long_bullets:  # nothing dropped
            self.assertIn(bullet.split(":")[0], text)

    def test_missing_image_file_renders_without_picture(self):
        slides = [slide(1, "Deck", "title"),
                  slide(2, "No image", bullets=["One: a", "Two: b", "Three: c"],
                        image_path=str(self.tmp / "missing.jpg"))]
        out = self.tmp / "missing.pptx"
        generate_pptx_with_officekit(self.deck(slides), out)
        self.assertEqual(check_pptx_layout(out)["pictures"], 0)

    def test_large_presentation(self):
        slides = [slide(1, "Big deck", "title")] + [
            slide(i, f"Topic {i}", bullets=[f"Item {i}.{j}: detail text" for j in range(4)],
                  paragraphs=[f"Paragraph for topic {i}."])
            for i in range(2, 41)
        ]
        out = self.tmp / "large.pptx"
        generate_pptx_with_officekit(self.deck(slides), out)
        report = check_pptx_layout(out)
        self.assertGreaterEqual(report["slides"], 40)
        self.assertEqual(report["issues"], [])

    def test_fallback_to_python_pptx_renderer(self):
        out = self.tmp / "fallback.pptx"
        with mock.patch("office.compono_pptx.render_deck", side_effect=RuntimeError("boom")):
            generate_pptx_with_officekit(self.deck(self.standard_slides()), out)
        names = {shape.name for s in PptPresentation(str(out)).slides for shape in s.shapes}
        self.assertNotIn("Process Step", names)  # legacy renderer was used
        self.assertTrue(out.is_file())

    def test_normalize_drops_model_reference_slides_and_schema_echo(self):
        data = {"title": "T", "slides": [
            {"title": "T", "layout": "title"},
            {"title": "Body", "bullets": ["Real point", "table", "image_query"]},
            {"title": "Sources", "bullets": ["Wikipedia: x"]},
        ]}
        result = PPTXContentGenerator._normalize(data, "t", "")
        self.assertEqual([s["title"] for s in result["slides"]], ["T", "Body"])
        self.assertEqual(result["slides"][1]["bullets"], ["Real point"])

    def test_normalize_keeps_title_slide_and_valid_citations(self):
        data = {"title": "Deck", "slides": [
            # The grammar forces bullets on the title slide too.
            {"title": "Deck", "layout": "title", "bullets": ["a", "b", "c"], "sources": [1]},
            {"title": "Body", "bullets": ["x: y"], "sources": [2, 2, 9, 0, "1"]},
        ]}
        result = PPTXContentGenerator._normalize(data, "t", "", source_count=3)
        self.assertEqual([s["title"] for s in result["slides"]], ["Deck", "Body"])  # no extra title slide
        self.assertEqual(result["slides"][0]["bullets"], [])
        self.assertEqual(result["slides"][1]["citations"], [2])  # only real, unique numbers

    def test_normalize_drops_half_empty_table(self):
        data = {"title": "Deck", "slides": [
            {"title": "Deck", "layout": "title"},
            {"title": "Benefits and Challenges", "layout": "data", "bullets": ["a: b"],
             "table": {"headers": ["Benefits", "Challenges"], "rows": [["Clean", ""], ["Dense", ""]]}},
        ]}
        result = PPTXContentGenerator._normalize(data, "t", "")
        self.assertIsNone(result["slides"][1]["table"])  # one filled column is not a table
        self.assertEqual(result["slides"][1]["layout"], "content")

    def fake_llm_generator(self, count):
        llm_slides = [{"title": "Deck", "subtitle": "s", "layout": "title", "paragraphs": [],
                       "bullets": ["a", "b", "c"], "table": None, "image_query": "q", "sources": []}]
        llm_slides += [{"title": f"Topic {i}", "subtitle": "s", "layout": "content",
                        "paragraphs": ["Text."], "bullets": ["k: v", "k2: v", "k3: v"],
                        "table": None, "image_query": "q", "sources": [1]} for i in range(2, count + 1)]
        generator = PPTXContentGenerator.__new__(PPTXContentGenerator)
        generator.llm = mock.Mock()
        generator.llm.generate_json.return_value = {"title": "Deck", "subtitle": "s", "slides": llm_slides}
        return generator

    def test_requested_slide_count_author_and_cited_notes(self):
        research = format_search_results(SOURCES)
        for count in (5, 8, 12):
            generator = self.fake_llm_generator(count)
            deck = generator.generate("t", research, slide_count=count, author="A. Writer")
            schema = generator.llm.generate_json.call_args.kwargs["schema"]
            self.assertEqual(schema["properties"]["slides"]["minItems"], count)
            self.assertEqual(schema["properties"]["slides"]["maxItems"], count)
            self.assertIn(f"Number of slides: {count}", generator.llm.generate_json.call_args.args[0])
            self.assertEqual(len(deck.slides), count)
            self.assertEqual(deck.author, "A. Writer")
            self.assertIn("[1] https://example.org/a", deck.slides[1].notes)
        with self.assertRaises(ValueError):
            self.fake_llm_generator(3).generate("t", research, slide_count=40)

    def test_author_citation_footer_and_numbered_references_in_every_engine(self):
        slides = self.standard_slides()
        slides[1]["citations"] = [1]
        deck = Presentation.model_validate({"title": "Energy Storage", "author": "A. Writer",
                                            "slides": slides, "sources": SOURCES})
        for engine in ("compono", "python-pptx", "both"):
            out = self.tmp / f"cite-{engine}.pptx"
            generate_pptx_with_officekit(deck, out, engine=engine)
            prs = PptPresentation(str(out))
            texts = [" ".join([sh.text_frame.text for sh in s.shapes if sh.has_text_frame]
                              + [c.text for sh in s.shapes if sh.has_table
                                 for row in sh.table.rows for c in row.cells])
                     for s in prs.slides]
            self.assertIn("A. Writer", texts[0], engine)
            self.assertIn("[1] Example source", " ".join(texts[-1:]), engine)  # references slide
            if engine != "compono":  # the footer text is drawn by python-pptx
                self.assertIn("Sources [1]", texts[1], engine)

    def test_long_title_still_validates(self):
        long_title = ("How grid scale battery energy storage systems reshape electricity markets "
                      "reliability and renewable integration across regions")
        slides = [slide(1, "Deck", "title"),
                  slide(2, long_title, bullets=["One: a", "Two: b", "Three: c"])]
        for engine in ("compono", "python-pptx", "both"):
            generate_pptx_with_officekit(self.deck(slides), self.tmp / f"long-{engine}.pptx", engine=engine)


class EngineOptionTests(unittest.TestCase):
    """The three terminal PPTX options + mandatory Office-Kit validation."""

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.img = make_image(self.tmp / "photo.jpg")
        self.deck = PptxTests.deck(self, PptxTests.standard_slides(self))

    def names(self, path):
        return {shape.name for s in PptPresentation(str(path)).slides for shape in s.shapes}

    def run_engine(self, engine):
        out = self.tmp / f"{engine}.pptx"
        with mock.patch.object(officekit_pptx_generator, "review_document",
                               wraps=review_document) as review:
            result = generate_pptx_with_officekit(self.deck, out, engine=engine)
        self.assertEqual(result, out)
        self.assertTrue(out.is_file())
        review.assert_called_once()  # validated once ...
        self.assertEqual(Path(review.call_args.args[0]), out)  # ... on the final file
        return out

    def test_option_1_compono_only(self):
        names = self.names(self.run_engine("compono"))
        self.assertNotIn("Accent Bar", names)  # no python-pptx brand pass
        self.assertNotIn("Process Step", names)

    def test_option_2_python_pptx(self):
        names = self.names(self.run_engine("python-pptx"))
        self.assertNotIn("Process Step", names)  # legacy renderer, no Compono sequence

    def test_option_3_both_in_one_file(self):
        out = self.run_engine("both")
        names = self.names(out)
        self.assertIn("Process Step", names)  # Compono sequence primitive ...
        self.assertIn("Accent Bar", names)    # ... restyled/added by python-pptx
        self.assertEqual(list(self.tmp.glob("*.pptx")), [out])  # no extra files

    def test_slide_count_input(self):
        self.assertEqual(main.parse_slide_count(""), 8)
        self.assertEqual(main.parse_slide_count(" 12 "), 12)
        for bad in ("0", "2", "99", "abc", "-5", "7.5"):
            with self.assertRaises(ValueError):
                main.parse_slide_count(bad)
        with mock.patch("builtins.input", side_effect=["x", "50", ""]), mock.patch("builtins.print"):
            self.assertEqual(main.ask_slide_count(), 8)

    def test_invalid_menu_selection_retries(self):
        with mock.patch("builtins.input", side_effect=["", "4", "abc", "3"]), \
                mock.patch("builtins.print") as printed:
            self.assertEqual(main.choose_pptx_engine(), "both")
        errors = [c for c in printed.call_args_list if "Invalid option" in str(c)]
        self.assertEqual(len(errors), 3)

    def test_explicit_engine_failure_does_not_fall_back(self):
        out = self.tmp / "fail.pptx"
        with mock.patch("office.compono_pptx.render_deck", side_effect=RuntimeError("boom")):
            with self.assertRaises(officekit_pptx_generator.PptxGenerationError) as ctx:
                generate_pptx_with_officekit(self.deck, out, engine="compono")
        self.assertIn("boom", str(ctx.exception))
        self.assertFalse(out.exists())

    def test_missing_output_file(self):
        out = self.tmp / "missing.pptx"
        fake = mock.Mock(return_value=self.tmp / "never_written.pptx")
        with mock.patch.dict(officekit_pptx_generator._RENDERERS, {"python-pptx": fake}):
            with self.assertRaises(officekit_pptx_generator.PptxGenerationError):
                generate_pptx_with_officekit(self.deck, out, engine="python-pptx")

    def test_invalid_output_path(self):
        with self.assertRaises(officekit_pptx_generator.PptxGenerationError):
            generate_pptx_with_officekit(self.deck, self.tmp / "deck.txt", engine="both")

    def test_officekit_validation_failure_keeps_file(self):
        out = self.run_engine("python-pptx")
        wrong = self.deck.model_dump()
        wrong["slides"][1]["title"] = "A title that is not in the file"
        with self.assertRaises(officekit_pptx_generator.PptxValidationError) as ctx:
            officekit_pptx_generator.validate_pptx(out, wrong)
        self.assertIn("A title that is not in the file", ctx.exception.review["missing_titles"])
        self.assertTrue(out.is_file())

    def test_officekit_exception_is_reported_not_swallowed(self):
        with mock.patch.object(officekit_pptx_generator, "review_document",
                               side_effect=OSError("officekit crashed")), \
                mock.patch.object(main, "PPTX_OUTPUT_DIR", self.tmp):
            result = main.render_with_officekit(self.deck, "crash topic", "both")
        self.assertIsNone(result)  # main reports failure, not success
        self.assertTrue((self.tmp / "crash_topic.pptx").is_file())  # kept for debugging


class ImageAndSearchTests(unittest.TestCase):

    def test_failed_download_returns_none(self):
        # Port 9 on localhost: connection refused immediately.
        with mock.patch.object(image_search, "cache_set"):
            self.assertIsNone(image_search.download_image("http://127.0.0.1:9/x.jpg",
                                                          Path(tempfile.mkdtemp())))

    def test_empty_image_search_results(self):
        deck = {"slides": [{"title": "Anything", "image_query": "solar battery farm", "layout": "content"}]}
        with mock.patch.object(image_search, "search_image_candidates", return_value=[]):
            image_search.enrich_presentation_with_images(deck, "topic")
        self.assertFalse(deck["slides"][0].get("image_path"))

    def test_empty_web_search_results(self):
        self.assertIn("No web search results", format_search_results([]))
        self.assertIn("No web search results", compact_research([]))
        text = compact_research([{"title": "A", "url": "https://x.org/very/long", "snippet": "b  c"}])
        self.assertEqual(text, "[1] A: b c")  # no URL tokens in the prompt


class DocxTests(unittest.TestCase):

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.report = Report.model_validate({
            "title": "Energy Storage Report",
            "subtitle": "Market and technology review",
            "executive_summary": "Storage is growing quickly. " * 10,
            "sections": [
                {"title": "Background", "paragraphs": ["Para one. " * 20],
                 "bullets": ["Grid: needs flexibility", "Solar: daytime peaks"],
                 "subsections": [{"title": "History", "paragraphs": ["Early pumped hydro. " * 8],
                                  "bullets": ["1890s: first plants"]}],
                 "image_path": make_image(self.tmp / "fig.jpg"), "image_source": "example.org",
                 "image_url": "https://example.org/fig.jpg", "image_caption": "A battery site"},
                {"title": "Market", "paragraphs": ["Para two. " * 20],
                 "table": {"headers": ["Year", "GW"], "rows": [["2022", "28"], ["2023", "45"]]},
                 "table_caption": "Installed capacity"},
            ],
            "conclusion": "Storage matters. " * 10,
            "sources": [{"title": "Source A", "url": "https://example.org/a"}],
        })

    def test_docx_structure(self):
        out = self.tmp / "report.docx"
        generate_docx_with_officekit(self.report, out)
        doc = Document(str(out))
        styles = [p.style.name for p in doc.paragraphs]
        for level in ("Heading 1", "Heading 2", "Heading 3", "Title", "List Bullet",
                      "List Number", "Caption"):
            self.assertIn(level, styles)
        self.assertEqual(len(doc.tables), 1)
        self.assertEqual(len(doc.inline_shapes), 1)
        self.assertEqual(doc.core_properties.title, "Energy Storage Report")
        xml = doc.element.xml
        self.assertIn("TOC \\o", xml)
        footer_xml = doc.sections[0].footer._element.xml
        self.assertIn("PAGE", footer_xml)
        self.assertIn("NUMPAGES", footer_xml)
        links = [r for r in doc.part.rels.values() if r.reltype.endswith("/hyperlink")]
        self.assertTrue(any(r.target_ref == "https://example.org/a" for r in links))
        review = review_document(out, expected_titles=["Background", "History", "Market"])
        self.assertTrue(review["valid"], review)

    def test_officekit_placeholder_repair(self):
        path = self.tmp / "placeholder.docx"
        doc = Document()
        doc.add_paragraph("Prepared for {{title}}")
        doc.save(str(path))
        review = review_document(path, placeholder_values={"title": "ACME"})
        self.assertEqual(review["edits_applied"], 1)
        self.assertTrue(review["valid"], review)
        self.assertIn("Prepared for ACME", Document(str(path)).paragraphs[0].text)


if __name__ == "__main__":
    unittest.main()
